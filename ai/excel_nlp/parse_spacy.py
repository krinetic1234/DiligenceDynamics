import openpyxl
import spacy
from openpyxl.chart import LineChart, Reference

nlp = spacy.load("en_core_web_sm")

def process_instruction(instruction):
    doc = nlp(instruction)
    numbers = [token.text for token in doc if token.pos_ == 'NUM']
    return numbers

def create_excel_list(numbers, sheet, start_row=1, start_column=1):
    for i, number in enumerate(numbers, start=start_row):
        sheet.cell(row=i, column=start_column, value=int(number))

def generate_graph(sheet, title, data_range):
    chart = LineChart()
    chart.title = title
    data = Reference(sheet, min_col=data_range['min_col'], min_row=data_range['min_row'], 
                     max_row=data_range['max_row'], max_col=data_range['max_col'])
    chart.add_data(data, titles_from_data=True)
    sheet.add_chart(chart, "E2")

def main(instruction):
    numbers = process_instruction(instruction)
    wb = openpyxl.Workbook()
    sheet = wb.active
    create_excel_list(numbers, sheet)
    generate_graph(sheet, "Stocks over Time", {'min_col': 1, 'min_row': 1, 'max_row': 4, 'max_col': 1})
    wb.save("SalesData.xlsx")

if __name__ == "__main__":
    instruction = "Apple recently has had fluctuating stock prices. Create a list with 80, 82, and 90."
    main(instruction)
